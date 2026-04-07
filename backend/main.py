import os
import logging
import zipfile
import tempfile
from datetime import datetime
from typing import List, Optional

from fastapi import FastAPI, HTTPException, Depends, Request, UploadFile, File, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from sqlalchemy import desc, or_

from . import models, schemas
from .database import SessionLocal, engine
from .services.diploma_generator import generar_diploma, generar_nombre_archivo
from .services.excel_processor import leer_excel, crear_plantilla_excel

logging.basicConfig(level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="DiplomasWeb — Sistema de Titulación", version="3.0.0")

app.add_middleware(CORSMiddleware, allow_origins=["*"],
    allow_credentials=False, allow_methods=["*"], allow_headers=["*"])

os.makedirs("diplomas_generados", exist_ok=True)
app.mount("/diplomas", StaticFiles(directory="diplomas_generados"), name="diplomas")

TEMPLATE_PATH = "templates/diploma_template.docx"
ZIP_TEMPORAL  = "diplomas_generados/_ultimo_zip.zip"

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    logger.error(f"Error no manejado: {exc}", exc_info=True)
    return JSONResponse(status_code=500, content={"detail": "Error interno del servidor."})


@app.get("/")
async def read_index():
    return FileResponse('frontend/index.html')

@app.post("/generar-diploma", tags=["Diplomas"])
async def crear_diploma(data: schemas.TitulacionCreate, db: Session = Depends(get_db)):
    existente = db.query(models.Titulacion).filter(models.Titulacion.dni == data.dni).first()
    if existente:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT,
            detail=f"Ya existe un diploma registrado para el DNI {data.dni}")
    try:
        db_titulacion = models.Titulacion(**data.dict())
        db.add(db_titulacion)
        db.commit()
        db.refresh(db_titulacion)
        nombre_base = generar_nombre_archivo(data.dict())
        output_base = f"diplomas_generados/{nombre_base}"
        ruta_docx, ruta_pdf = generar_diploma(data.dict(), TEMPLATE_PATH, output_base)
        db_titulacion.ruta_documento = ruta_docx
        db_titulacion.ruta_pdf       = ruta_pdf
        db_titulacion.updated_at     = datetime.now()
        db.commit()
        logger.info(f"Diploma generado para DNI: {data.dni}")
        return {"mensaje": f"Diploma generado para {data.nombre} {data.apellidos}",
                "dni": data.dni, "id": db_titulacion.id, "tiene_pdf": ruta_pdf is not None}
    except FileNotFoundError as e:
        db.rollback()
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/validar/{dni}", tags=["Validación"])
async def validar_diploma(dni: str, db: Session = Depends(get_db)):
    import re
    if not re.match(r"^\d{8}$", dni):
        raise HTTPException(status_code=400, detail="DNI inválido")
    t = db.query(models.Titulacion).filter(models.Titulacion.dni == dni).first()
    if not t:
        raise HTTPException(status_code=404, detail="No se encontró diploma para este DNI")
    return {"valido": True, "nombre_completo": f"{t.nombre} {t.apellidos}",
            "dni": t.dni, "carrera": t.carrera, "modalidad": t.modalidad,
            "titulo_trabajo": t.titulo_trabajo, "asesor": t.asesor,
            "fecha": str(t.fecha), "mensaje": "✅ Diploma AUTÉNTICO verificado por el sistema."}

@app.get("/admin/diplomas", response_model=List[schemas.TitulacionResponse], tags=["Admin"])
async def listar_diplomas(skip: int = 0, limit: int = 200,
    buscar: Optional[str] = None, modalidad: Optional[str] = None,
    db: Session = Depends(get_db)):
    query = db.query(models.Titulacion)
    if buscar:
        b = f"%{buscar.lower()}%"
        query = query.filter(or_(
            models.Titulacion.nombre.ilike(b), models.Titulacion.apellidos.ilike(b),
            models.Titulacion.dni.ilike(b), models.Titulacion.carrera.ilike(b),
            models.Titulacion.numero_resolucion.ilike(b)))
    if modalidad:
        query = query.filter(models.Titulacion.modalidad == modalidad)
    return query.order_by(desc(models.Titulacion.created_at)).offset(skip).limit(limit).all()

@app.get("/admin/diplomas/{id}", response_model=schemas.TitulacionResponse, tags=["Admin"])
async def obtener_diploma(id: int, db: Session = Depends(get_db)):
    d = db.query(models.Titulacion).filter(models.Titulacion.id == id).first()
    if not d:
        raise HTTPException(status_code=404, detail="Diploma no encontrado")
    return d

@app.put("/admin/diplomas/{id}", response_model=schemas.TitulacionResponse, tags=["Admin"])
async def editar_diploma(id: int, data: schemas.TitulacionUpdate, db: Session = Depends(get_db)):
    d = db.query(models.Titulacion).filter(models.Titulacion.id == id).first()
    if not d:
        raise HTTPException(status_code=404, detail="Diploma no encontrado")
    for field, value in data.dict(exclude_unset=True).items():
        setattr(d, field, value)
    d.updated_at = datetime.now()
    db.commit()
    db.refresh(d)
    return d

@app.delete("/admin/diplomas/{id}", tags=["Admin"])
async def eliminar_diploma(id: int, db: Session = Depends(get_db)):
    d = db.query(models.Titulacion).filter(models.Titulacion.id == id).first()
    if not d:
        raise HTTPException(status_code=404, detail="Diploma no encontrado")
    for ruta in [d.ruta_documento, d.ruta_pdf]:
        if ruta and os.path.exists(ruta):
            os.remove(ruta)
    db.delete(d)
    db.commit()
    return {"mensaje": f"Diploma de {d.nombre} {d.apellidos} eliminado"}

@app.get("/admin/estadisticas", tags=["Admin"])
async def estadisticas(db: Session = Depends(get_db)):
    from sqlalchemy import func
    total    = db.query(models.Titulacion).count()
    tesis    = db.query(models.Titulacion).filter(models.Titulacion.modalidad == "Tesis").count()
    tsp      = db.query(models.Titulacion).filter(models.Titulacion.modalidad == "TSP").count()
    carreras = db.query(models.Titulacion.carrera,
                func.count(models.Titulacion.id).label("total")
               ).group_by(models.Titulacion.carrera).all()
    return {"total_diplomas": total, "por_modalidad": {"Tesis": tesis, "TSP": tsp},
            "por_carrera": [{"carrera": c, "total": t} for c, t in carreras]}

@app.get("/admin/descargar/{dni}/{formato}", tags=["Admin"])
async def descargar_diploma(dni: str, formato: str, db: Session = Depends(get_db)):
    t = db.query(models.Titulacion).filter(models.Titulacion.dni == dni).first()
    if not t:
        raise HTTPException(status_code=404, detail="Diploma no encontrado")
    if formato == "docx":
        ruta  = t.ruta_documento
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ext   = "docx"
    elif formato == "pdf":
        ruta  = t.ruta_pdf
        media = "application/pdf"
        ext   = "pdf"
    else:
        raise HTTPException(status_code=400, detail="Formato inválido. Usa 'docx' o 'pdf'")
    if not ruta or not os.path.exists(ruta):
        raise HTTPException(status_code=404, detail=f"Archivo {formato.upper()} no encontrado")
    return FileResponse(path=ruta, filename=f"diploma_{t.nombre}_{t.apellidos}.{ext}", media_type=media)
@app.get("/admin/descargar-seleccion", tags=["Admin"])
async def descargar_seleccion(
    buscar: Optional[str] = None,
    modalidad: Optional[str] = None,
    carrera: Optional[str] = None,
    fecha_desde: Optional[str] = None,
    fecha_hasta: Optional[str] = None,
    resolucion: Optional[str] = None,
    db: Session = Depends(get_db)
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from datetime import date as date_type

    query = db.query(models.Titulacion)

    if buscar:
        b = f"%{buscar.lower()}%"
        query = query.filter(or_(
            models.Titulacion.nombre.ilike(b),
            models.Titulacion.apellidos.ilike(b),
            models.Titulacion.dni.ilike(b),
        ))
    if modalidad:
        query = query.filter(models.Titulacion.modalidad == modalidad)
    if carrera:
        query = query.filter(models.Titulacion.carrera.ilike(f"%{carrera}%"))
    if resolucion:
        query = query.filter(
            models.Titulacion.numero_resolucion.ilike(f"%{resolucion}%"))
    if fecha_desde:
        try:
            query = query.filter(
                models.Titulacion.fecha >= date_type.fromisoformat(fecha_desde))
        except: pass
    if fecha_hasta:
        try:
            query = query.filter(
                models.Titulacion.fecha <= date_type.fromisoformat(fecha_hasta))
        except: pass

    diplomas = query.order_by(desc(models.Titulacion.created_at)).all()
    if not diplomas:
        raise HTTPException(status_code=404,
            detail="No hay registros con los filtros seleccionados")

    # Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Selección"
    headers = ["ID","Resolución","Nombre","Apellidos","DNI",
               "Carrera","Modalidad","Título","Asesor","Fecha"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F1F3D",
                                end_color="0F1F3D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 20
    for row, d in enumerate(diplomas, 2):
        ws.append([d.id, d.numero_resolucion or "", d.nombre, d.apellidos,
                   d.dni, d.carrera, d.modalidad, d.titulo_trabajo,
                   d.asesor, str(d.fecha)])

    excel_path = "diplomas_generados/_seleccion_temp.xlsx"
    wb.save(excel_path)

    # ZIP
    zip_path = "diplomas_generados/_seleccion.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.write(excel_path, "lista_seleccion.xlsx")
        for d in diplomas:
            if d.ruta_documento and os.path.exists(d.ruta_documento):
                zf.write(d.ruta_documento, os.path.basename(d.ruta_documento))
            if d.ruta_pdf and os.path.exists(d.ruta_pdf):
                zf.write(d.ruta_pdf, os.path.basename(d.ruta_pdf))

    total = len(diplomas)
    nombre_zip = f"diplomas_seleccion_{total}registros.zip"
    return FileResponse(path=zip_path, filename=nombre_zip,
                        media_type="application/zip")

@app.get("/admin/exportar-excel", tags=["Admin"])
async def exportar_excel(db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    diplomas = db.query(models.Titulacion).order_by(desc(models.Titulacion.created_at)).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diplomas"
    headers = ["ID","Resolución","Nombre","Apellidos","DNI","Carrera","Modalidad","Título","Asesor","Fecha","Registrado"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F1F3D", end_color="0F1F3D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 18
    for row, d in enumerate(diplomas, 2):
        ws.append([d.id, d.numero_resolucion or "", d.nombre, d.apellidos, d.dni,
                   d.carrera, d.modalidad, d.titulo_trabajo, d.asesor,
                   str(d.fecha), str(d.created_at)[:19] if d.created_at else ""])
    path = "diplomas_generados/_exportacion.xlsx"
    wb.save(path)
    return FileResponse(path=path, filename="diplomas_exportados.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/admin/plantilla-excel", tags=["Admin"])
async def descargar_plantilla():
    path = "diplomas_generados/_plantilla.xlsx"
    crear_plantilla_excel(path)
    return FileResponse(path=path, filename="plantilla_carga_masiva.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/admin/carga-masiva", tags=["Admin"])
async def carga_masiva(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx o .xls")
    ruta_temporal = os.path.join(tempfile.gettempdir(), f"upload_{os.getpid()}.xlsx")
    with open(ruta_temporal, "wb") as f_tmp:
        f_tmp.write(await file.read())
    try:
        filas, errores = leer_excel(ruta_temporal)
    finally:
        try:
            os.unlink(ruta_temporal)
        except Exception:
            pass
    if not filas and errores:
        raise HTTPException(status_code=400, detail="; ".join(errores))
    exitosos, omitidos, advertencias = 0, 0, list(errores)
    archivos_generados = []
    for fila in filas:
        try:
            existente = db.query(models.Titulacion).filter(
                models.Titulacion.dni == str(fila["dni"])).first()
            if existente:
                advertencias.append(f"DNI {fila['dni']} ya existe — omitido")
                omitidos += 1
                continue
            db_t = models.Titulacion(
                numero_resolucion=fila.get("numero_resolucion"),
                nombre=str(fila["nombre"]), apellidos=str(fila["apellidos"]),
                dni=str(fila["dni"]), carrera=str(fila["carrera"]),
                modalidad=str(fila["modalidad"]), titulo_trabajo=str(fila["titulo_trabajo"]),
                asesor=str(fila["asesor"]), fecha=fila["fecha"])
            db.add(db_t)
            db.commit()
            db.refresh(db_t)
            nombre_base = generar_nombre_archivo(fila)
            output_base = f"diplomas_generados/{nombre_base}"
            ruta_docx, ruta_pdf = generar_diploma(fila, TEMPLATE_PATH, output_base)
            db_t.ruta_documento = ruta_docx
            db_t.ruta_pdf       = ruta_pdf
            db_t.updated_at     = datetime.now()
            db.commit()
            archivos_generados.append(ruta_docx)
            if ruta_pdf:
                archivos_generados.append(ruta_pdf)
            exitosos += 1
        except Exception as e:
            db.rollback()
            advertencias.append(f"Error en fila DNI {fila.get('dni','?')}: {str(e)}")
    tiene_pdf = any(f.endswith(".pdf") for f in archivos_generados)
    if archivos_generados:
        with zipfile.ZipFile(ZIP_TEMPORAL, "w") as zf:
            for archivo in archivos_generados:
                if os.path.exists(archivo):
                    zf.write(archivo, os.path.basename(archivo))
    return {"procesados": len(filas), "exitosos": exitosos,
            "omitidos": omitidos, "advertencias": advertencias, "tiene_pdf": tiene_pdf}

@app.get("/admin/descargar-zip", tags=["Admin"])
async def descargar_zip():
    if not os.path.exists(ZIP_TEMPORAL):
        raise HTTPException(status_code=404, detail="No hay ZIP disponible")
    return FileResponse(path=ZIP_TEMPORAL, filename="diplomas_generados.zip",
                        media_type="application/zip")
@app.get("/health")
def health_check():
    return {"status": "healthy"}

from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse

app.mount("/static", StaticFiles(directory="frontend"), name="static")

@app.get("/")
async def index():
    return FileResponse("frontend/index.html")

@app.get("/admin")
async def admin():
    return FileResponse("frontend/admin.html")

@app.get("/verificar")
async def verificar():
    return FileResponse("frontend/verificar.html")

@app.get("/carga-masiva")
async def carga():
    return FileResponse("frontend/carga-masiva.html")