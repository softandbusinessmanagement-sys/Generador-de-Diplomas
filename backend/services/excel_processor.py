import openpyxl
import os
import logging
from datetime import date, datetime
from typing import List, Tuple

logger = logging.getLogger(__name__)

COLUMNAS_REQUERIDAS = ["nombre","apellidos","dni","carrera","modalidad","titulo_trabajo","asesor","fecha"]


def leer_excel(file_path: str) -> Tuple[List[dict], List[str]]:
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
    except Exception as e:
        return [], [f"No se pudo abrir el Excel: {e}"]

    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in headers]
    if faltantes:
        return [], [f"Columnas faltantes: {', '.join(faltantes)}"]

    filas, errores = [], []

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(c is None for c in row):
            continue

        fila = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        errs = []

        for col in COLUMNAS_REQUERIDAS:
            if not fila.get(col) and fila.get(col) != 0:
                errs.append(f"'{col}' vacío")

        dni = str(fila.get("dni","")).strip().split(".")[0]
        if dni and (not dni.isdigit() or len(dni) != 8):
            errs.append(f"DNI inválido: '{dni}'")

        if str(fila.get("modalidad","")).strip() not in ["Tesis","TSP"]:
            errs.append(f"Modalidad inválida (usar Tesis o TSP)")

        fecha_raw = fila.get("fecha")
        if fecha_raw:
            if isinstance(fecha_raw, (date, datetime)):
                fila["fecha"] = fecha_raw.date() if isinstance(fecha_raw, datetime) else fecha_raw
            else:
                for fmt in ["%Y-%m-%d", "%d/%m/%Y"]:
                    try:
                        fila["fecha"] = datetime.strptime(str(fecha_raw).strip(), fmt).date()
                        break
                    except:
                        pass
                else:
                    errs.append(f"Fecha inválida: '{fecha_raw}'")

        if errs:
            errores.append(f"Fila {row_num}: {', '.join(errs)}")
            continue

        for k in fila:
            if isinstance(fila[k], str):
                fila[k] = fila[k].strip()
        fila["dni"] = str(fila["dni"]).strip().split(".")[0]

        filas.append(fila)

    return filas, errores


def crear_plantilla_excel(output_path: str):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Titulaciones"

    headers = ["numero_resolucion","nombre","apellidos","dni","carrera","modalidad","titulo_trabajo","asesor","fecha"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F1F3D", end_color="0F1F3D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 22

    ejemplo = ["RES-2025-001","Juan Carlos","García López","12345678",
               "Ingeniería de Sistemas","Tesis","Sistema de Gestión con IA","Dr. María Pérez Torres","2025-06-15"]
    for col, v in enumerate(ejemplo, 1):
        ws.cell(row=2, column=col, value=v)

    ws.cell(row=3, column=1, value="← Elimina esta fila de ejemplo")
    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path