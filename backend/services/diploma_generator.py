
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import qrcode
import os
import logging
import subprocess
import shutil

from dotenv import load_dotenv
load_dotenv()

logger = logging.getLogger(__name__)

MESES = {
    1:"enero", 2:"febrero", 3:"marzo", 4:"abril",
    5:"mayo", 6:"junio", 7:"julio", 8:"agosto",
    9:"septiembre", 10:"octubre", 11:"noviembre", 12:"diciembre"
}


def reemplazar_en_runs(paragraph, marcador: str, valor: str):
    texto = "".join(r.text for r in paragraph.runs)
    if marcador in texto:
        nuevo = texto.replace(marcador, valor)
        for i, run in enumerate(paragraph.runs):
            run.text = nuevo if i == 0 else ""
        return True
    return False


def preparar_data(data: dict) -> dict:
    data = data.copy()
    fecha = data.get("fecha")
    if hasattr(fecha, "day"):
        data["fecha_formato"] = f"{fecha.day} de {MESES[fecha.month]} de {fecha.year}"
    else:
        data["fecha_formato"] = str(fecha)
    data["nombre_completo"] = f"{data.get('nombre','')} {data.get('apellidos','')}"
    if not data.get("numero_resolucion"):
        data["numero_resolucion"] = "S/N"
    return data


def generar_docx(data: dict, template_path: str, output_path: str) -> str:
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    data = preparar_data(data)
    qr_path = None

    try:
        doc = Document(template_path)

        for paragraph in doc.paragraphs:
            for key, value in data.items():
                reemplazar_en_runs(paragraph, f"{{{{{key}}}}}", str(value) if value else "")

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in data.items():
                            reemplazar_en_runs(paragraph, f"{{{{{key}}}}}", str(value) if value else "")

        # QR
        base_url = os.getenv("BASE_URL", "http://127.0.0.1:5500/frontend")
        qr_url = f"{base_url}/verificar.html?dni={data['dni']}"
        qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_H, box_size=10, border=4)
        qr.add_data(qr_url)
        qr.make(fit=True)
        qr_img = qr.make_image(fill_color="black", back_color="white")

        qr_path = f"diplomas_generados/qr_{data['dni']}.png"
        os.makedirs("diplomas_generados", exist_ok=True)
        qr_img.save(qr_path)

        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(qr_path, width=Inches(1.5))

        cap = doc.add_paragraph("Escanea para verificar autenticidad")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(8)

        doc.save(output_path)
        return output_path

    finally:
        if qr_path and os.path.exists(qr_path):
            os.remove(qr_path)


def docx_a_pdf(docx_path: str, pdf_path: str):
    try:
        output_dir = os.path.dirname(pdf_path)
        libreoffice_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "soffice", "libreoffice"
        ]
        soffice = None
        for path in libreoffice_paths:
            if os.path.exists(path) or shutil.which(path):
                soffice = path
                break

        if not soffice:
            logger.warning("LibreOffice no encontrado. Instálalo para generar PDF.")
            return None

        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", output_dir, docx_path],
            capture_output=True, timeout=30
        )
        base = os.path.splitext(os.path.basename(docx_path))[0]
        generated = os.path.join(output_dir, f"{base}.pdf")
        if os.path.exists(generated):
            if generated != pdf_path:
                os.rename(generated, pdf_path)
            return pdf_path
        return None
    except Exception as e:
        logger.error(f"Error PDF: {e}")
        return None


def generar_nombre_archivo(data: dict) -> str:
    resolucion = str(data.get("numero_resolucion", "")).strip()
    dni = data.get("dni", "")
    if resolucion and resolucion != "S/N":
        limpia = "".join(c for c in resolucion if c.isalnum() or c in "-_")
        return f"diploma_{limpia}_{dni}"
    return f"diploma_{dni}"


def generar_diploma(data: dict, template_path: str, output_base: str):
    """Retorna (ruta_docx, ruta_pdf)"""
    ruta_docx = generar_docx(data, template_path, f"{output_base}.docx")
    ruta_pdf  = docx_a_pdf(ruta_docx, f"{output_base}.pdf")
    return ruta_docx, ruta_pdf